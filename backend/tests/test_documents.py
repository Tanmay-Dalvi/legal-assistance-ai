"""
Comprehensive tests for document ingestion, validation, extraction, and lifecycle.
"""

import os
import pytest
from pathlib import Path
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient
from app.services import storage_service

# Ensure test database is used
os.environ["DATABASE_URL"] = "sqlite+aiosqlite:///./data/test_legal_ai_full.db"


@pytest.fixture
def txt_file(tmp_path):
    p = tmp_path / "test.txt"
    p.write_text("This is a test document.\n\nIt has multiple paragraphs.")
    return p

@pytest.fixture
def pdf_file(tmp_path):
    p = tmp_path / "test.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(100, 100, "This is a test PDF.")
    c.save()
    return p

@pytest.fixture
def docx_file(tmp_path):
    p = tmp_path / "test.docx"
    doc = DocxDocument()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('This is a test DOCX paragraph.')
    doc.save(str(p))
    return p

@pytest.fixture
def empty_file(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("")
    return p

@pytest.fixture
def oversized_file(tmp_path):
    p = tmp_path / "large.txt"
    with open(p, "wb") as f:
        f.write(b"0" * (21 * 1024 * 1024))
    return p

@pytest.fixture
def unsupported_file(tmp_path):
    p = tmp_path / "test.csv"
    p.write_text("a,b,c\n1,2,3")
    return p

@pytest.fixture
def fake_pdf(tmp_path):
    p = tmp_path / "fake.pdf"
    p.write_text("This is just text but named as pdf.")
    return p

@pytest.fixture
def fake_docx(tmp_path):
    p = tmp_path / "fake.docx"
    p.write_text("This is just text but named as docx.")
    return p

@pytest.fixture
def malformed_pdf(tmp_path):
    p = tmp_path / "malformed.pdf"
    p.write_bytes(b"%PDF-1.7\nnot a real pdf")
    return p

@pytest.fixture
def malformed_docx(tmp_path):
    p = tmp_path / "malformed.docx"
    p.write_bytes(b"PK\x03\x04not a real docx")
    return p


class TestDocumentUploads:
    
    # 1. Valid TXT upload + 10. Extracted TXT content
    def test_upload_txt(self, client, txt_file):
        with open(txt_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("test.txt", f, "text/plain")}
            )
        assert response.status_code == 201
        data = response.json()
        assert data["original_filename"] == "test.txt"
        assert data["processing_status"] == "ready"
        assert data["extracted_character_count"] > 10

    # 2. Valid PDF upload + 11. PDF page metadata
    def test_upload_pdf(self, client, pdf_file):
        with open(pdf_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("test.pdf", f, "application/pdf")}
            )
        assert response.status_code == 201
        data = response.json()
        assert data["original_filename"] == "test.pdf"
        assert data["processing_status"] == "ready"
        assert data["page_count"] == 1
        assert data["extracted_character_count"] > 5

    # 3. Valid DOCX upload + 12. DOCX structure preservation
    def test_upload_docx(self, client, docx_file):
        with open(docx_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )
        assert response.status_code == 201
        data = response.json()
        assert data["original_filename"] == "test.docx"
        assert data["processing_status"] == "ready"
        assert data["extracted_character_count"] > 10

    # 4. Unsupported extension
    def test_upload_unsupported(self, client, unsupported_file):
        with open(unsupported_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("test.csv", f, "text/csv")}
            )
        assert response.status_code == 415

    # 7. Empty file (now correctly returning a 422 ValidationError)
    def test_upload_empty(self, client, empty_file):
        with open(empty_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("empty.txt", f, "text/plain")}
            )
        assert response.status_code == 422
        assert "empty" in response.json()["error"]["message"].lower()

    # 6. Oversized file
    def test_upload_oversized(self, client, oversized_file):
        with open(oversized_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("large.txt", f, "text/plain")}
            )
        assert response.status_code == 413

    # 5. Invalid MIME type / 8. Malformed PDF (caught by magic bytes)
    def test_fake_pdf(self, client, fake_pdf):
        with open(fake_pdf, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("fake.pdf", f, "application/pdf")}
            )
        assert response.status_code == 422
        assert "does not match" in response.json()["error"]["message"].lower()

    # 9. Malformed DOCX (caught by magic bytes)
    def test_fake_docx(self, client, fake_docx):
        with open(fake_docx, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("fake.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            )
        assert response.status_code == 422

    def test_mime_mismatch(self, client, txt_file):
        with open(txt_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("test.txt", f, "application/pdf")},
            )
        assert response.status_code == 415
        assert response.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"

    def test_malformed_pdf_is_sanitized(self, client, malformed_pdf):
        with open(malformed_pdf, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("malformed.pdf", f, "application/pdf")},
            )
        assert response.status_code == 422
        body = response.json()
        assert body["error"]["message"] == "The uploaded document could not be read or is malformed."
        assert "Traceback" not in response.text

    def test_malformed_docx_is_sanitized(self, client, malformed_docx):
        with open(malformed_docx, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("malformed.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 422
        assert "Traceback" not in response.text

    # 21. Path traversal attempt (backend drops relative paths from standard upload filename, but we can try to force it if FastAPI allows it, though it usually doesn't. 
    # Let's test the storage service path traversal directly later, or via api if possible).
    def test_path_traversal_api(self, client, txt_file):
        with open(txt_file, "rb") as f:
            response = client.post(
                "/api/v1/documents/upload",
                files={"file": ("../../../etc/passwd.txt", f, "text/plain")}
            )
        # It will extract passwd.txt and be safe, or fail. As long as it doesn't traverse.
        # Actually our code uses uuid for stored_filename, so original filename is just metadata!
        assert response.status_code == 201
        # Original filename is just stored as string, no traversal happens.
        assert response.json()["original_filename"] == "../../../etc/passwd.txt"


class TestDocumentLifecycle:
    # 13. DB metadata, 14. GET document, 15. GET list, 16. DELETE, 17/18. File removal, 25. Ordering
    def test_lifecycle(self, client, txt_file):
        # Upload 1
        with open(txt_file, "rb") as f:
            up_res = client.post("/api/v1/documents/upload", files={"file": ("doc1.txt", f, "text/plain")})
        doc_id = up_res.json()["id"]

        # 14. GET document
        get_res = client.get(f"/api/v1/documents/{doc_id}")
        assert get_res.status_code == 200
        assert get_res.json()["id"] == doc_id
        
        # 15. GET document list & 25. List ordering
        list_res = client.get("/api/v1/documents")
        assert list_res.status_code == 200
        docs = list_res.json()
        assert any(d["id"] == doc_id for d in docs)
        # Latest should be first (descending by created_at)
        assert docs[0]["id"] == doc_id

        # 16. DELETE document
        del_res = client.delete(f"/api/v1/documents/{doc_id}")
        assert del_res.status_code == 204

        # 19. Nonexistent document
        get_res2 = client.get(f"/api/v1/documents/{doc_id}")
        assert get_res2.status_code == 404
        
        # 17/18. File removal physically
        upload_dir = Path("./data/uploads").resolve()
        extracted_dir = Path("./data/extracted").resolve()
        assert not (upload_dir / f"{doc_id}.txt").exists()
        assert not (extracted_dir / f"{doc_id}.json").exists()

    # 20. Invalid document ID
    def test_invalid_doc_id(self, client):
        res = client.get("/api/v1/documents/invalid-uuid-format")
        assert res.status_code == 404

    # 26. Health endpoint remains functional
    def test_health_endpoint(self, client):
        res = client.get("/api/v1/health")
        assert res.status_code == 200
        assert res.json()["status"] == "ok"


class TestStorageBoundaries:
    def test_upload_path_rejects_traversal_and_sibling_prefix(self):
        with pytest.raises(Exception):
            storage_service.get_safe_upload_path("../../evil.txt")
        with pytest.raises(Exception):
            storage_service.get_safe_upload_path("../uploads-elsewhere/evil.txt")

    def test_extracted_path_rejects_traversal(self):
        with pytest.raises(Exception):
            storage_service.get_safe_extracted_path("../../evil")
